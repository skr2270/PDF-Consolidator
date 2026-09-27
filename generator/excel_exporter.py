"""OpenPyXL 3-Tab Warehouse Unloading Checklist Exporter."""
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from consolidator.item_grouper import ConsignmentSummary

class ExcelExporter:
    """Generates 3-tab warehouse unloading workbook."""

    HEADER_FILL = PatternFill(start_color="3C3D3A", end_color="3C3D3A", fill_type="solid")
    ACCENT_FILL = PatternFill(start_color="EAEAEA", end_color="EAEAEA", fill_type="solid")
    WHITE_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    BOLD_FONT = Font(name="Calibri", size=11, bold=True)
    NORMAL_FONT = Font(name="Calibri", size=10)

    @classmethod
    def export(cls, summary: ConsignmentSummary, output_path: str):
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        wb = openpyxl.Workbook()

        # TAB 1: Manifest Summary
        ws1 = wb.active
        ws1.title = "Manifest Summary"
        ws1.views.sheetView[0].showGridLines = True
        
        ws1.append(["CONSIGNMENT DISPATCH MANIFEST"])
        ws1.append([])
        ws1.append(["Dispatch Date:", summary.dispatch_date])
        ws1.append(["Source Store:", summary.store_name])
        ws1.append(["Destination Store:", summary.destination])
        ws1.append(["Total Bags / Orders:", summary.total_bags])
        ws1.append(["Total Consolidated SKUs:", summary.total_skus])
        ws1.append(["Total Pieces Dispatched:", summary.total_pieces])
        ws1.append(["Total Consignment Value (INR):", summary.total_value])
        
        ws1["A1"].font = Font(name="Calibri", size=14, bold=True)
        for r in range(3, 10):
            ws1[f"A{r}"].font = cls.BOLD_FONT
            ws1[f"B{r}"].font = cls.NORMAL_FONT
        ws1["B9"].number_format = "#,##0.00"

        # TAB 2: Consolidated SKUs
        ws2 = wb.create_sheet(title="Consolidated SKUs")
        ws2.views.sheetView[0].showGridLines = True
        headers_tab2 = ["#", "SKU", "Item Description", "Transfer Price (INR)", "Quantity (Pcs)", "Subtotal (INR)", "Contributing Bags Count", "Bags List"]
        ws2.append(headers_tab2)
        for cell in ws2[1]:
            cell.fill = cls.HEADER_FILL
            cell.font = cls.WHITE_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for idx, it in enumerate(summary.items, 1):
            ws2.append([
                idx,
                it.sku,
                it.item_name,
                it.cost_price,
                it.quantity,
                it.subtotal,
                len(it.bags),
                ", ".join(it.bags)
            ])
            r = idx + 1
            ws2[f"A{r}"].alignment = Alignment(horizontal="center")
            ws2[f"D{r}"].number_format = "#,##0.00"
            ws2[f"E{r}"].alignment = Alignment(horizontal="right")
            ws2[f"F{r}"].number_format = "#,##0.00"

        # TAB 3: Bag-by-Bag Unloading Checklist
        ws3 = wb.create_sheet(title="Bag Unloading Checklist")
        ws3.views.sheetView[0].showGridLines = True
        headers_tab3 = ["Bag #", "Order Number", "Total Items in Bag", "Total Quantity (Pcs)", "Order Value (INR)", "Physical Seal Verified", "Received Sign"]
        ws3.append(headers_tab3)
        for cell in ws3[1]:
            cell.fill = cls.HEADER_FILL
            cell.font = cls.WHITE_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for idx, o in enumerate(summary.orders, 1):
            ws3.append([
                f"Bag {idx:02d}",
                o.metadata.order_number,
                len(o.items),
                o.metadata.total_quantity,
                o.metadata.grand_total,
                "[  ] Verified",
                ""
            ])
            r = idx + 1
            ws3[f"A{r}"].alignment = Alignment(horizontal="center")
            ws3[f"D{r}"].alignment = Alignment(horizontal="right")
            ws3[f"E{r}"].number_format = "#,##0.00"
            ws3[f"F{r}"].alignment = Alignment(horizontal="center")

        for sheet in wb.worksheets:
            for col in sheet.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = col[0].column_letter
                sheet.column_dimensions[col_letter].width = max(max_len + 3, 12)

        wb.save(output_path)
        return output_path
